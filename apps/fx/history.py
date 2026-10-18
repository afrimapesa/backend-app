"""Daily mid-market rate history (for trend charts) and bulk import of historical rates."""

from __future__ import annotations

import csv
import re
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.db import connection

from .services import PIVOT

# --- reading history -------------------------------------------------------------------

_DAILY_SQL = """
    SELECT time_bucket('1 day', observed_at) AS day, last(rate, observed_at) AS rate
    FROM fx_mid_rates
    WHERE base_id = %s AND quote_id = %s AND observed_at >= %s AND observed_at < %s
    GROUP BY day
    ORDER BY day
"""


@dataclass(frozen=True)
class DailySeries:
    points: list[tuple[date, Decimal]]
    derived: bool  # True when computed as a cross rate through USD


def daily_history(base: str, quote: str, *, days: int, today: date) -> DailySeries:
    """Last observed rate per UTC day for 1 ``base`` in ``quote``, oldest first.

    Falls back to a USD cross rate (USD/quote / USD/base) on days where both legs exist.
    """
    start = datetime.combine(today - timedelta(days=days - 1), time.min, tzinfo=UTC)
    end = datetime.combine(today + timedelta(days=1), time.min, tzinfo=UTC)

    direct = _daily(base, quote, start, end)
    if direct or PIVOT in (base, quote):
        return DailySeries(points=sorted(direct.items()), derived=False)

    usd_base = _daily(PIVOT, base, start, end)
    usd_quote = _daily(PIVOT, quote, start, end)
    crossed = {day: usd_quote[day] / usd_base[day] for day in usd_base.keys() & usd_quote.keys()}
    return DailySeries(points=sorted(crossed.items()), derived=bool(crossed))


def _daily(base: str, quote: str, start: datetime, end: datetime) -> dict[date, Decimal]:
    with connection.cursor() as cursor:
        cursor.execute(_DAILY_SQL, [base, quote, start, end])
        return {day.date(): rate for day, rate in cursor.fetchall()}


# --- importing history -----------------------------------------------------------------

PAIR = re.compile(r"^\s*([A-Za-z]{3})\s*[/\-_ ]?\s*([A-Za-z]{3})\s*$")
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y")


@dataclass(frozen=True)
class HistoricalRate:
    day: date
    base: str
    quote: str
    rate: Decimal

    @property
    def observed_at(self) -> datetime:
        return datetime.combine(self.day, time.min, tzinfo=UTC)


class ImportFormatError(ValueError):
    pass


def read_rows(path: Path) -> list[list[object]]:
    """Read a CSV or Excel (.xlsx) file into rows of cell values (first sheet for Excel)."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return [list(row) for row in csv.reader(f)]
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    raise ImportFormatError(f"Unsupported file type {suffix!r}; use .csv or .xlsx")


def parse_rates(rows: list[list[object]]) -> Iterator[HistoricalRate | str]:
    """Yield parsed rates, or an error string for each row/cell that can't be read.

    Accepts either layout:

    * long: columns ``date, base, quote, rate`` (any order; ``pair`` like "USD/NGN"
      may replace base + quote);
    * wide: a ``date`` column plus one column per pair, headed "USD/NGN", "USDNGN", ...
    """
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise ImportFormatError("File is empty")
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    if "date" not in header:
        raise ImportFormatError("No 'date' column found in the first row")
    date_col = header.index("date")

    if "rate" in header:
        yield from _parse_long(header, rows[1:], date_col)
    else:
        yield from _parse_wide(rows[0], rows[1:], date_col)


def _parse_long(
    header: list[str], body: Iterable[list[object]], date_col: int
) -> Iterator[HistoricalRate | str]:
    rate_col = header.index("rate")
    pair_col = header.index("pair") if "pair" in header else None
    base_col = header.index("base") if "base" in header else None
    quote_col = header.index("quote") if "quote" in header else None
    if pair_col is None and (base_col is None or quote_col is None):
        raise ImportFormatError("Long format needs 'base' and 'quote' columns, or 'pair'")
    for n, row in enumerate(body, start=2):
        try:
            if pair_col is not None:
                base, quote = _pair(row[pair_col])
            else:
                assert base_col is not None and quote_col is not None
                base, quote = (
                    str(row[base_col]).strip().upper(),
                    str(row[quote_col]).strip().upper(),
                )
            yield HistoricalRate(_date(row[date_col]), base, quote, _rate(row[rate_col]))
        except (ValueError, IndexError) as exc:
            yield f"row {n}: {exc}"


def _parse_wide(
    raw_header: list[object], body: Iterable[list[object]], date_col: int
) -> Iterator[HistoricalRate | str]:
    pairs: dict[int, tuple[str, str]] = {}
    for i, cell in enumerate(raw_header):
        if i == date_col or cell in (None, ""):
            continue
        try:
            pairs[i] = _pair(cell)
        except ValueError:
            yield f"column {i + 1}: header {cell!r} is not a currency pair; skipped"
    if not pairs:
        raise ImportFormatError("Wide format needs currency-pair column headers like 'USD/NGN'")
    for n, row in enumerate(body, start=2):
        try:
            day = _date(row[date_col])
        except (ValueError, IndexError) as exc:
            yield f"row {n}: {exc}"
            continue
        for i, (base, quote) in pairs.items():
            cell = row[i] if i < len(row) else None
            if cell in (None, ""):
                continue  # no observation that day for this pair
            try:
                yield HistoricalRate(day, base, quote, _rate(cell))
            except ValueError as exc:
                yield f"row {n}, {base}/{quote}: {exc}"


def _pair(value: object) -> tuple[str, str]:
    match = PAIR.match(str(value or ""))
    if not match:
        raise ValueError(f"{value!r} is not a currency pair")
    return match.group(1).upper(), match.group(2).upper()


def _date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognised date {value!r}")


def _rate(value: object) -> Decimal:
    try:
        rate = Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        raise ValueError(f"rate {value!r} is not a number") from None
    if not rate.is_finite() or rate <= 0:
        raise ValueError(f"rate {value!r} must be a positive number")
    return rate
